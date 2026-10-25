import os
from docx import Document

SRC = "Diagramas finales/prueba2_apa7_v2.docx"

def inspect_figma_section():
    if not os.path.exists(SRC):
        return

    doc = Document(SRC)
    start_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if "Figura 32" in p.text:
            start_idx = i
            break
    
    if start_idx == -1:
        print("Figura 32 not found")
        return

    print(f"--- Paragraphs after Figura 32 (index {start_idx}) ---")
    for i in range(start_idx, min(start_idx + 25, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        text = p.text.strip()
        has_image = any("drawing" in r._element.xml or "picture" in r._element.xml for r in p.runs)
        img_mark = "[IMAGE]" if has_image else ""
        print(f"[{i}] {img_mark} {text}")

if __name__ == "__main__":
    inspect_figma_section()

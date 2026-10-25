import os
import re
from docx import Document

SRC = "Diagramas finales/prueba2_apa7_v2.docx"
DST = "Diagramas finales/prueba2_apa7_v2.docx"

def fix_and_add_figure():
    doc = Document(SRC)
    
    # 1. Fix the malformed Figura 48 at 827
    # Current: [827] |Figura 48 Pruebas funcionales: creación de cuestionarios con IA|
    # We want to split it.
    p48 = doc.paragraphs[827]
    if "Figura 48" in p48.text:
        full_text = p48.text
        p48.text = "Figura 48"
        p48.style = "Figure Label"
        
        # Insert Title
        p_title = doc.add_paragraph()
        p48._element.addnext(p_title._element)
        p_title.style = "Caption"
        title_text = full_text.replace("Figura 48", "").strip()
        if not title_text:
            title_text = "Pruebas Funcionales: Creación de Cuestionarios con IA (Paso 3)"
        rt = p_title.add_run(title_text)
        rt.italic = True

    # 2. Renumber figures from 49 upwards (+1)
    # We do this BEFORE inserting the new 49 to avoid double renumbering
    for p in doc.paragraphs:
        match = re.search(r'Figura (\d+)', p.text)
        if match:
            num = int(match.group(1))
            if num >= 49:
                for r in p.runs:
                    if f"Figura {num}" in r.text:
                        r.text = r.text.replace(f"Figura {num}", f"Figura {num+1}")
                if f"Figura {num}" in p.text and f"Figura {num+1}" not in p.text:
                     p.text = p.text.replace(f"Figura {num}", f"Figura {num+1}")

    # 3. Find the unlabeled image at index 839 (it might have shifted slightly due to split at 827)
    # Splitting 827 added 1 paragraph, so 839 is now 840.
    target_idx = 840
    target_p = doc.paragraphs[target_idx]
    
    # Verify it's the right one (empty text, has image)
    if not target_p.text.strip() and any("drawing" in r._element.xml or "picture" in r._element.xml for r in target_p.runs):
        target_p.style = "Figure Label"
        target_p.text = "Figura 49"
        
        p_title_new = doc.add_paragraph()
        target_p._element.addnext(p_title_new._element)
        p_title_new.style = "Caption"
        rtn = p_title_new.add_run("Pruebas Funcionales: Generación de Cuestionarios con IA")
        rtn.italic = True
        
        p_note_new = doc.add_paragraph()
        p_title_new._element.addnext(p_note_new._element)
        p_note_new.style = "Figure Note"
        p_note_new.add_run("Nota. pruebas funcionales: generación de cuestionarios con IA, por , 2026.")
    else:
        print(f"Warning: Paragraph at {target_idx} is not what we expected: '{target_p.text}'")

    doc.save(DST)
    print("Figura 48 fixed, new Figura 49 added, and subsequent figures renumbered.")

if __name__ == "__main__":
    fix_and_add_figure()

# -*- coding: utf-8 -*-
"""Second pass on `CorreccionesF.docx`:
  - Re-write every purple (correction) paragraph so it contains ONLY the full
    corrected text (no commentary, no parenthetical justification, no "no se usa X").
  - Leave the orange (explanation) paragraphs untouched — they are the only place
    where the rationale lives.
  - Re-affirm legend wording.
  - Idempotent: re-running it is safe.
The script does not delete or insert paragraphs; it just rewrites the runs of
already-existing purple paragraphs so the format is preserved.
"""
import shutil
from docx import Document
from docx.shared import RGBColor, Pt

SRC = r"documentos proyecto/FORMATO PARA ENTREGA DE DOCUMENTO FINAL - 2024 3 CorreccionesF.docx"

RED_HEX_SET = {'C62828', 'FF0000', 'C00000', 'E60000', 'BF0000', 'CC0000'}
PURPLE_HEX_SET = {'6A1B9A', '7030A0', '800080', '660066', '7B2882', '9B59B6', '993399', '7E22CE'}

PURPLE = RGBColor(0x6A, 0x1B, 0x9A)

# Backup before mutating
shutil.copyfile(SRC, SRC.replace('.docx', '_backup2.docx'))

doc = Document(SRC)


def first_run_color(p):
    for r in p.runs:
        c = r.font.color and r.font.color.rgb
        if c is None:
            continue
        s = str(c).upper()
        if s in RED_HEX_SET:
            return 'R'
        if s in PURPLE_HEX_SET:
            return 'P'
    return None


def rewrite_purple_text(p, clean_text):
    """Replace the runs of a purple paragraph with a single purple run carrying clean_text."""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(clean_text)
    run.font.size = Pt(11)
    run.font.color.rgb = PURPLE


# (red prefix → cleaned purple full text).  The match is by a stable
# substring of the red paragraph so reordering doesn't break the mapping.
CLEAN_BY_RED_SUBSTR = [
    (
        "Figura 3. Diagrama de clases UML",
        "✏️ Figura 3. Diagrama de clases UML: módulo Auth (hexagonal + clean architecture). "
        "Fuente: documentos proyecto/diagrams/drawIO/05_Diagrama_Clases_UML.drawio.",
    ),
    (
        "Figura 5. Diagrama de secuencia",
        "✏️ Figura 5. Diagrama de secuencia: inicio de sesión OAuth con Auth.js v5 + bridge a NestJS. "
        "Fuente: documentos proyecto/diagrams/drawIO/07_Diagrama_Secuencias.drawio.",
    ),
    (
        "Figura 6. Diagrama de secuencia",
        "✏️ Figura 6. Diagrama de secuencia: resolución de un quiz e integración con FastAPI. "
        "Fuente: documentos proyecto/diagrams/drawIO/07_Diagrama_Secuencias.drawio.",
    ),
    (
        "El componente ModelFactory",
        "✏️ El componente model_factory se encarga de la selección dinámica del modelo correspondiente "
        "de acuerdo con el tipo de pregunta y la configuración utilizada. Los resultados de clasificación se "
        "persisten gracias al adaptador GradingRepo, implementado mediante SQLAlchemy asíncrono, mientras las "
        "evaluaciones por lotes se gestionan mediante concurrencia asíncrona con asyncio.gather.",
    ),
    (
        "Observer: procesamiento de eventos asíncronos mediante Celery",
        "✏️ Observer: notificación de eventos de dominio mediante EventEmitter2 en NestJS.",
    ),
    (
        "comunicación entre ambos servicios tuvo lugar mediante medio de servicios web de tipo REST",
        "✏️ La comunicación entre ambos servicios se realiza mediante servicios web de tipo REST utilizando "
        "como soporte el modo de transporte seguro Transport Layer Security en formato JavaScript Object Notation. Se "
        "implementó un mecanismo de autenticación entre servicios para generar interacciones seguras mediante el "
        "envío en el encabezado Authorization Bearer de JSON Web Token firmados digitalmente.",
    ),
    (
        "petrolero de IA crea evaluaciones",
        "✏️ El motor de IA crea evaluaciones en un promedio de 8 segundos.",
    ),
    (
        "Implementar calificación automática con FastAPI y la entrega en la retroalimentación humanizada del petrolero",
        "✏️ Objetivo: Implementar la calificación automática con FastAPI y la entrega de "
        "retroalimentación humanizada por el agente de IA.",
    ),
    (
        "Verificar la comunicación REST y gRPC entre los microservicios",
        "✏️ Verificar la comunicación REST entre los microservicios desarrollados en NestJS y FastAPI para la "
        "calificación IA y el agente.",
    ),
    (
        "Pruebas de integración: revisar la interacción entre módulos",
        "✏️ Pruebas de integración: revisar la interacción entre módulos, microservicios, bases de "
        "datos y APIs REST, garantizando que la interacción entre los distintos componentes sea la correcta.",
    ),
    (
        "Validaciones de la comunicación entre microservicios mediante protocolos REST y gR",
        "✏️ Validaciones de la comunicación entre microservicios mediante el protocolo REST.",
    ),
]


def find_clean_for(red_text):
    for substr, clean in CLEAN_BY_RED_SUBSTR:
        if substr in red_text:
            return clean
    return None


# Walk the doc, find every red paragraph (after the legend), then find the next
# purple paragraph within a few lines, and rewrite its content.
paragraphs = doc.paragraphs
changed = 0
for i, p in enumerate(paragraphs):
    if i < 6:
        # legend zone — skip
        continue
    if first_run_color(p) != 'R':
        continue
    clean = find_clean_for(p.text)
    if not clean:
        continue
    # Locate the next purple paragraph within the next 5 lines
    for j in range(i + 1, min(i + 6, len(paragraphs))):
        if first_run_color(paragraphs[j]) == 'P':
            rewrite_purple_text(paragraphs[j], clean)
            changed += 1
            break

# Also re-affirm legend wording (idempotent rewrite).
LEGEND_RED = paragraphs[2]
LEGEND_PURPLE = paragraphs[3]


def rewrite_with_color(p, text, color):
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = color


rewrite_with_color(
    LEGEND_RED,
    "Rojo + ❌: texto del documento que NO cumple con lo que ofrece el aplicativo PROYDEMO. Se conserva tal cual para evidenciar el hallazgo.",
    RGBColor(0xC0, 0x00, 0x00),
)
rewrite_with_color(
    LEGEND_PURPLE,
    "Morado + ✏️: corrección con el TEXTO COMPLETO que sí corresponde al aplicativo PROYDEMO (sin explicaciones; reemplazo recomendado).",
    PURPLE,
)

doc.save(SRC)
print("OK guardado:", SRC)
print("Purple paragraphs limpiados:", changed)

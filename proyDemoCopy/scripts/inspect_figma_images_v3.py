import os
from docx import Document

SRC = "Diagramas finales/prueba2_apa7_v2.docx"

def inspect_figma_section():
    if not os.path.exists(SRC):
        return

    doc = Document(SRC)
    start_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if "Prototipos de Interfaz: Figma" in p.text:
            start_idx = i
            break
    
    if start_idx == -1:
        print("Section not found")
        return

    print(f"--- Paragraphs in Figma section (from index {start_idx}) ---")
    # Search for Figura tags
    for i in range(start_idx, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        text = p.text.strip()
        has_image = any("drawing" in r._element.xml or "picture" in r._element.xml for r in p.runs)
        if "Figura" in text or has_image:
             img_mark = "[IMAGE]" if has_image else ""
             print(f"[{i}] {img_mark} {text}")
        
        if "Agente IA: Arquitectura" in text: # End of section usually
            break

if __name__ == "__main__":
    inspect_figma_section()

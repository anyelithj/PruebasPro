from docx import Document

doc = Document("Diagramas finales/prueba2_apa7_v2.docx")
start_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "Figura 48" in p.text:
        start_idx = i
        break

if start_idx == -1:
    print("Figura 48 not found")
else:
    print(f"--- Inspection from index {start_idx} ---")
    for i in range(start_idx, min(start_idx + 20, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        has_image = any("drawing" in r._element.xml or "picture" in r._element.xml for r in p.runs)
        img_mark = "[IMAGE]" if has_image else ""
        print(f"[{i}] {img_mark} Style: {p.style.name:15} | Text: '{p.text}'")

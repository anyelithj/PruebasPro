"""Append section 10 to Auth_y_Docker.docx:
 - Por que NO se usa MCP (Model Context Protocol)
 - Por que NO se usa API Gateway dedicado
 - Tipos de notificaciones implementadas en ProyDemo

Idempotente: si la seccion 10 ya existe, la reescribe.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

DOCX_PATH = Path(
    r"C:/Users/angui/Documents/proyDemo/documentos proyecto/Auth_y_Docker.docx"
)

MARKER_H1 = "10. Decisiones de arquitectura: MCP, API Gateway y Notificaciones"


def remove_existing_section(doc: Document) -> None:
    body = doc.element.body
    paragraphs = list(doc.paragraphs)
    start_idx = None
    for i, p in enumerate(paragraphs):
        if p.text.strip() == MARKER_H1:
            start_idx = i
            break
    if start_idx is None:
        return
    for p in paragraphs[start_idx:]:
        body.remove(p._element)


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def main() -> None:
    doc = Document(str(DOCX_PATH))

    # quitar sentinel final
    body = doc.element.body
    for p in doc.paragraphs[-3:]:
        if "Fin del documento" in p.text:
            body.remove(p._element)
            break

    remove_existing_section(doc)

    doc.add_heading(MARKER_H1, level=1)

    # 10.1 Por que no MCP
    doc.add_heading("10.1 Por que NO se usa MCP (Model Context Protocol)", level=2)
    doc.add_paragraph(
        "MCP (Model Context Protocol) es un protocolo estandar para conectar modelos "
        "de lenguaje con herramientas, fuentes de datos y servidores de contexto externos. "
        "En ProyDemo se evaluo y se descarto su uso por las siguientes razones:"
    )
    bullet(
        doc,
        "Caso de uso acotado: la generacion de preguntas y el feedback academico se "
        "resuelven con una sola llamada de inferencia a modelos locales (Qwen3:8b / "
        "Mistral Nemo) servidos por Ollama. No hay necesidad de orquestar multiples "
        "herramientas ni de exponer contexto dinamico al modelo, que es el problema que "
        "MCP resuelve.",
    )
    bullet(
        doc,
        "Sin agentes ni tool-calling: el sistema no implementa agentes autonomos que "
        "decidan que herramienta invocar. El flujo es deterministico (prompt -> JSON de "
        "preguntas), por lo que la capa de protocolo de MCP seria sobre-ingenieria.",
    )
    bullet(
        doc,
        "Costo de integracion vs beneficio: adoptar MCP exigiria levantar y mantener "
        "servidores MCP, definir su ciclo de vida y seguridad, sin aportar valor sobre "
        "la integracion REST directa con Ollama que ya cubre el requerimiento.",
    )
    bullet(
        doc,
        "Extensibilidad ya cubierta por diseno: el patron Factory + Strategy + Chain of "
        "Responsibility (AIProviderFactory) permite agregar nuevos proveedores de IA sin "
        "refactor mayor. La extensibilidad buscada no requiere MCP.",
    )
    bullet(
        doc,
        "Conclusion: MCP queda documentado como posible evolucion futura (por ejemplo si "
        "se incorporan agentes con acceso a multiples fuentes), pero hoy no se usa porque "
        "incrementaria la complejidad sin un beneficio funcional inmediato.",
    )

    # 10.2 Por que no API Gateway dedicado
    doc.add_heading("10.2 Por que NO se usa un API Gateway dedicado", level=2)
    doc.add_paragraph(
        "ProyDemo no incorpora un API Gateway dedicado (Kong, Tyk, Apigee, AWS API "
        "Gateway). En su lugar usa el patron Backend-For-Frontend (BFF) con los Route "
        "Handlers de Next.js. Las razones:"
    )
    bullet(
        doc,
        "Topologia simple: el sistema tiene pocos servicios (Next.js, NestJS, FastAPI, "
        "Ollama). Un gateway dedicado se justifica cuando hay muchos microservicios y "
        "necesidad de enrutamiento, versionado y agregacion centralizada; aqui seria "
        "infraestructura adicional sin retorno.",
    )
    bullet(
        doc,
        "Responsabilidades ya distribuidas: autenticacion (JwtAuthGuard global), rate "
        "limiting (CustomThrottlerGuard), normalizacion de errores (HttpExceptionFilter) "
        "y logging/timeout/transform (interceptors globales) ya estan implementados en "
        "NestJS. El BFF de Next.js agrega el proxy y el manejo de cookies httpOnly.",
    )
    bullet(
        doc,
        "Seguridad del token: el BFF mantiene el access/refresh token en cookies "
        "httpOnly del lado servidor y los inyecta como Bearer hacia NestJS, evitando "
        "exponerlos al navegador. Esto cumple el objetivo de un gateway (no filtrar "
        "credenciales) sin un componente extra.",
    )
    bullet(
        doc,
        "Costo operativo: un gateway dedicado anade un punto de fallo, latencia y "
        "configuracion que no aportan valor en el alcance academico actual.",
    )
    bullet(
        doc,
        "Conclusion: el patron BFF cubre hoy las funciones de un gateway. Si el sistema "
        "crece a una malla de microservicios, se podria introducir un gateway dedicado "
        "sin reescribir la logica de negocio.",
    )

    # 10.3 Tipos de notificaciones
    doc.add_heading("10.3 Tipos de notificaciones implementadas", level=2)
    doc.add_paragraph(
        "El modulo Notifications define un catalogo cerrado de tipos en el enum "
        "NotificationType (entidad Notification, persistida en la tabla 'notifications' "
        "con soporte de marcado como leida). Los tipos contemplados son:"
    )
    bullet(doc, "ACHIEVEMENT: logro o insignia desbloqueada por el estudiante.")
    bullet(doc, "SCORE: resultado o puntaje obtenido tras completar un examen.")
    bullet(doc, "EMAIL_VERIFY: verificacion de correo electronico (flujo de registro).")
    bullet(doc, "PASSWORD_RESET: solicitud de restablecimiento de contrasena.")
    bullet(doc, "QUIZ_PUBLISHED: aviso de que un nuevo examen fue publicado.")
    bullet(doc, "LOGIN_DEVICE: inicio de sesion desde un dispositivo (alerta de seguridad).")
    bullet(doc, "SYSTEM: mensajes generales del sistema.")
    doc.add_paragraph(
        "Canales de entrega: el correo transaccional (verificacion de email y "
        "restablecimiento de contrasena) se entrega via SMTP, capturado en desarrollo por "
        "MailHog (puertos 1025 SMTP / 8025 UI, ver docker-compose.dev.yml). La estructura "
        "para notificaciones realtime via SSE (Server-Sent Events) esta prevista en el "
        "modulo (sse.controller) y la persistencia para historial offline ya existe en la "
        "entidad Notification."
    )
    doc.add_paragraph(
        "Estado actual: el catalogo de tipos y la entidad persistente estan definidos; el "
        "correo de verificacion y reset estan operativos. El resto de tipos quedan "
        "preparados para activarse mediante el caso de uso de broadcast y el canal SSE."
    )

    p = doc.add_paragraph("- Fin del documento -")
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.save(str(DOCX_PATH))
    print("OK: seccion 10 agregada")


if __name__ == "__main__":
    main()

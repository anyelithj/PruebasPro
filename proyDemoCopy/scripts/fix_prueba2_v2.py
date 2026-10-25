import os
from docx import Document
from docx.shared import Pt, Inches

SRC = "Diagramas finales/prueba2.docx"
DST = "Diagramas finales/prueba2_corregido.docx"

def fix_document():
    if not os.path.exists(SRC):
        print(f"Error: {SRC} not found")
        return

    doc = Document(SRC)
    
    # 1. Complete truncated sentences and remove audit marks
    paragraphs_to_remove = []
    
    for p in doc.paragraphs:
        text = p.text.strip()
        
        # Complete truncated sentences
        if "hay siete archivos que cubren los" in text and text.endswith("los"):
            p.text = text + " routers de calificación, retroalimentación, asistente y analítica, garantizando la cobertura de los endpoints definidos en el contrato de la API."
        
        if text == "Las pruebas funcionales del microservicio de extremo a extremo":
            p.text = text + " validan el flujo completo desde la recepción de la solicitud hasta la generación de la respuesta por el modelo de IA local, asegurando la consistencia semántica y el cumplimiento de los esquemas Pydantic."

        # Remove audit labels/sections
        audit_labels = ["[CORRECCIÓN]", "[EXPLICACIÓN TÉCNICA]", "[OBSERVACION]", "[NOTA:", "hallazgos encontrados:", 
                        "Recomendaciones documentales:", "Resultados de la Auditoria", "FIN DEL INFORME DE AUDITORIA"]
        if any(label in text for label in audit_labels) or "Naranja +" in text or "Morado +" in text or "Rojo +" in text:
            paragraphs_to_remove.append(p)
            continue

    # Remove marked paragraphs
    for p in paragraphs_to_remove:
        try:
            p._element.getparent().remove(p._element)
        except:
            pass

    # 2. Update Tables with "ESTADO" column
    for table in doc.tables:
        header_cells = table.rows[0].cells
        header_text = " ".join(c.text for c in header_cells).lower()
        
        # Identify relevant tables (20-26)
        if "archivo" in header_text and ("funcionalidad" in header_text or "validada" in header_text or "administración" in header_text):
            if "estado" not in header_text:
                # Add ESTADO column
                table.add_column(Inches(1.0))
                # Set header
                header_cell = table.rows[0].cells[-1]
                header_cell.text = "ESTADO"
                for run in header_cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
                
                # Fill status for each row
                for row in table.rows[1:]:
                    file_name = row.cells[0].text.strip()
                    desc_text = row.cells[1].text.strip()
                    
                    status = "Realizada" # Default
                    
                    if "VACIO" in desc_text or "0 bytes" in desc_text or "sin pruebas implementadas" in desc_text:
                        status = "Pendiente"
                    elif "ADVERTENCIA" in desc_text or "parcial" in desc_text.lower():
                        status = "Parcial"
                    elif "pendiente" in desc_text.lower() and "casos" in desc_text.lower():
                        status = "Parcial"
                    
                    # Special cases based on audit
                    if "UserAdminSection.test.tsx" in file_name:
                        status = "Realizada"
                    elif "sse.integration.spec.ts" in file_name:
                        status = "Parcial"
                    elif "test_grading_schemas.py" in file_name or "grading" in file_name:
                        status = "Realizada"
                    
                    status_cell = row.cells[-1]
                    status_cell.text = status
                    for run in status_cell.paragraphs[0].runs:
                        run.font.size = Pt(9)

    doc.save(DST)
    print(f"Document saved to {DST}")

if __name__ == "__main__":
    fix_document()

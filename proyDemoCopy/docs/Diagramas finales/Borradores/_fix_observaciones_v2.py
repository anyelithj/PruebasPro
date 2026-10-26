# -*- coding: utf-8 -*-
"""Corrige observaciones especificas del usuario sobre prueba.docx."""
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from copy import deepcopy

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"

RED    = RGBColor(0xFF, 0x00, 0x00)
PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
GREEN  = RGBColor(0x00, 0x80, 0x00)

doc = Document(SRC)
log = []

def new_para_after(para, text, color):
    el = OxmlElement('w:p')
    para._element.addnext(el)
    np = type(para)(el, doc)
    r = np.add_run(text)
    r.font.color.rgb = color
    return np

def paint_run(run, color):
    run.font.color.rgb = color

def paint_para(para, color):
    for run in para.runs:
        run.font.color.rgb = color

# ================================================================
# Helper: find paragraph by exact text prefix
# ================================================================
def find_para(text_start):
    for p in doc.paragraphs:
        if p.text.strip().startswith(text_start):
            return p
    return None

# ================================================================
# FIX 1: P210 PURPLE - solo entidades completas
# ================================================================
p210 = find_para('[CORRECCI')
if p210 and 'ocho entidades' in p210.text:
    for run in p210.runs:
        if 'ocho entidades' in run.text or 'Adicionalmente' in run.text:
            run.text = ("Se han determinado las siguientes entidades completamente implementadas en el modelo: "
                       "Users, AuthSessions, EmailVerificationTokens, PasswordResetTokens, Quizzes, Questions, "
                       "QuizSessions, Scores y FeedbackIA (persistencia de retroalimentaci\u00f3n).")
            log.append("P210: PURPLE corregido - solo entidades completas")
            break

# ================================================================
# FIX 2: P211 NARANJA - incluir ambas entidades incompletas
# ================================================================
p211 = find_para('[EXPLICACI')
if p211:
    for run in p211.runs:
        if 'Achievement' in run.text:
            run.text = (
                "Las entidades Achievement y Notification existen como tablas TypeORM pero no est\u00e1n completamente implementadas: "
                "Achievements carece de API REST y frontend (0 referencias en Frontend/src/); "
                "Notifications tiene el controlador SSE vac\u00edo (0 l\u00edneas) y casos de uso sin implementar. "
                "Los diagramas ER y Relacional las excluyen como 'incompletas'.")
            log.append("P211: ORANGE actualizado con ambas entidades")
            break

# ================================================================
# FIX 3: P212 [NOTA] de VERDE a NARANJA
# ================================================================
p212 = find_para('[NOTA]')
if p212 and 'Achievement' in p212.text:
    paint_para(p212, ORANGE)
    log.append("P212: [NOTA] cambiado de VERDE a NARANJA")

# ================================================================
# FIX 4: P550 - mantener VERDE original, agregar correccion
# ================================================================
p550 = find_para('Gestion de Usuarios')
if p550 and 'HU12' in p550.text:
    # Insert ORANGE note after
    new_para_after(p550,
        "[OBSERVACI\u00d3N] El texto menciona 'HU12 (retroalimentaci\u00f3n personalizada IA)' pero HU12 no existe en el backlog del producto (solo HU01-HU11 y HU14). "
        "La retroalimentaci\u00f3n personalizada con IA corresponde a la HU08. La gesti\u00f3n de usuarios implementa la HU14 (CRUD completo via /admin/users).",
        ORANGE)
    log.append("P550: observacion NARANJA agregada (HU12 no existe)")

# ================================================================
# FIX 5: P553 - mantener VERDE original, agregar correccion
# ================================================================
p553 = find_para('Gestion Global de Evaluaciones')
if p553 and 'HU12' in p553.text:
    new_para_after(p553,
        "[OBSERVACI\u00d3N] El texto menciona 'HU12 (retroalimentaci\u00f3n personalizada IA)' pero HU12 no existe en el backlog del producto. "
        "La retroalimentaci\u00f3n personalizada con IA corresponde a la HU08. La gesti\u00f3n global de evaluaciones permite moderar el cat\u00e1logo pero no implementa retroalimentaci\u00f3n IA.",
        ORANGE)
    log.append("P553: observacion NARANJA agregada (HU12 no existe)")

# ================================================================
# FIX 6: Actualizar seccion auditoria - hallazgo #1
# ================================================================
p_audit = find_para('1) P\u00e1rrafo')
if p_audit:
    for run in p_audit.runs:
        if 'Correcci' in run.text or 'ocho entidades' in run.text:
            run.text = (
                "1) P\u00e1rrafo 'ocho entidades... incluye Achievements' marcado ROJO. "
                "La correcci\u00f3n MORADO lista solo las entidades completamente implementadas "
                "(Users, AuthSessions, EmailVerificationTokens, PasswordResetTokens, Quizzes, "
                "Questions, QuizSessions, Scores, FeedbackIA). Achievements y Notification "
                "se explican en el texto NARANJA como entidades existentes pero incompletas.")
            log.append("Auditoria: hallazgo #1 actualizado")
            break

# ================================================================
# SAVE
# ================================================================
doc.save(SRC)
print(f"Guardado: {SRC}")
print(f"\nCambios ({len(log)}):")
for l in log:
    print(f"  {l}")

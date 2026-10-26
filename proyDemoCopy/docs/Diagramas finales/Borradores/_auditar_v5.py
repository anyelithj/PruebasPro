# -*- coding: utf-8 -*-
"""Audita Prueba.docx - version final y precisa.
CORRECCIONES:
1. Marca ROJO + agrega MORADO para afirmaciones incorrectas desactualizadas
2. Agrega AZUL para funcionalidades existentes en codigo NO documentadas
3. Marca ROJO filas de pruebas vacias (Tabla 21 y 22 del documento)
4. Agrega seccion Resultados de la Auditoria al final
"""
import shutil
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"
BACKUP = f"{BASE}/prueba_backup_final.docx"
shutil.copy(SRC, BACKUP)

RED    = RGBColor(0xFF, 0x00, 0x00)
BLUE   = RGBColor(0x00, 0x00, 0xFF)
PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
BLACK  = RGBColor(0x00, 0x00, 0x00)

doc = Document(SRC)
log = []

def add_after(para, text, color=None, bold=False):
    el = OxmlElement('w:p')
    para._element.addnext(el)
    np = type(para)(el, doc)
    r = np.add_run(text)
    if color: r.font.color.rgb = color
    if bold: r.font.bold = True
    return np

def paint_para(para, color):
    for run in para.runs:
        run.font.color.rgb = color

def paint_row(table, row_idx, color):
    for cell in table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = color

# ================================================================
# 1. NOTIFICACIONES - heading incorrecto
# ================================================================
for p in doc.paragraphs:
    if p.text.strip() == "Notificaciones en tiempo real y centro de notificaciones":
        paint_para(p, RED)
        log.append("ROJO: Notificaciones en tiempo real")
        add_after(p,
            "[EXPLICACI\u00d3N T\u00c9CNICA] El m\u00f3dulo de notificaciones tiene entidad pero los casos de uso, controlador SSE y gateway est\u00e1n vac\u00edos.",
            ORANGE)
        add_after(p,
            "[CORRECCI\u00d3N] Notificaciones en tiempo real (propuesta): la entidad Notification existe en el dominio pero los casos de uso y controlador SSE carecen de implementaci\u00f3n funcional.",
            PURPLE)
        break

# ================================================================
# 2. CORREGIR afirmacion desactualizada: "Edicion de perfil: pagina en construccion"
#    Esta afirmacion es INCORRECTA porque la pagina de edicion de perfil YA existe.
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if "Edicion de perfil de usuario" in txt and "construccion" in txt.lower():
        paint_para(p, RED)
        log.append(f"ROJO: afirmacion desactualizada sobre perfil: {txt[:60]}...")
        add_after(p,
            "[EXPLICACI\u00d3N T\u00c9CNICA] La p\u00e1gina de edici\u00f3n de perfil (Frontend/src/app/(portal)/profile/edit/page.tsx, 101 l\u00edneas) y el endpoint PATCH /profile/me (backend, 78 l\u00edneas) est\u00e1n implementados y funcionales.",
            ORANGE)
        add_after(p,
            "[CORRECCI\u00d3N] Edici\u00f3n de perfil de usuario: funcionalidad completa. El endpoint PATCH /profile/me permite actualizar nombre, secci\u00f3n y datos del perfil. La p\u00e1gina /profile/edit en el frontend proporciona el formulario de edici\u00f3n correspondiente.",
            PURPLE)
        break

# ================================================================
# 3. AGREGAR AZUL: funcionalidades implementadas NO documentadas
# ================================================================
# 3a: Perfil publico despues de descripcion de perfil
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "perfil" in txt and "visualizaci" in txt and "progreso" in txt:
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] Perfil p\u00fablico: el endpoint GET /profile/[username] y la p\u00e1gina /profile/[username] exponen las estad\u00edsticas de cualquier usuario (puntajes, evaluaciones completadas, racha de actividad) sin requerir autenticaci\u00f3n, permitiendo compartir el progreso acad\u00e9mico.",
            BLUE)
        log.append("AZUL: perfil publico agregado")
        break

# 3b: Historial de estudiantes para docente
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "resultados de sus estudiantes" in txt or "resultados de los estudiantes" in txt:
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] Vista detallada por estudiante: los docentes pueden acceder a /teacher/students/[id]/history para ver el historial completo de un estudiante con an\u00e1lisis por categor\u00eda, dificultad, tipo de pregunta, tendencias de rendimiento y recomendaciones del sistema.",
            BLUE)
        log.append("AZUL: historial detallado estudiante agregado")
        break

# 3c: Reportes de administrador
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "administrador" in txt and ("gestion" in txt or "gesti\u00f3n" in txt) and "usuario" in txt:
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] Reportes de administrador: la p\u00e1gina /admin/reports presenta estad\u00edsticas globales (usuarios por rol, evaluaciones creadas, quizzes por estado, distribuci\u00f3n de puntajes) con gr\u00e1ficos y tablas exportables.",
            BLUE)
        log.append("AZUL: reportes admin agregados")
        break

# 3d: Caching del leaderboard
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "clasificaci" in txt and ("leaderboard" in txt or "ranking" in txt or "mejores" in txt):
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] Cach\u00e9 del leaderboard: el endpoint GET /score/leaderboard utiliza CacheModule con Redis (TTL 60 segundos) para optimizar el rendimiento en consultas concurrentes.",
            BLUE)
        log.append("AZUL: cache leaderboard agregado")
        break

# 3e: Weakness analysis IA
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "retroalimentaci" in txt and ("personalizada" in txt or "agente" in txt) and "ia" in txt:
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] An\u00e1lisis de debilidades: el endpoint POST /ai/feedback/weaknesses y el componente WeaknessChart identifican las \u00e1reas de bajo rendimiento del estudiante, generando recomendaciones pedag\u00f3gicas espec\u00edficas basadas en sus respuestas.",
            BLUE)
        log.append("AZUL: weakness analysis agregado")
        break

# 3f: Tutor IA asistente
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip().lower()
    if "asistente" in txt and "inteligencia artificial" in txt:
        add_after(p,
            "[FUNCIONALIDAD ADICIONAL] Tutor IA interactivo: la p\u00e1gina /assistant y el componente AssistantChat permiten al estudiante realizar consultas pedag\u00f3gicas ad-hoc al agente de IA, recibiendo explicaciones detalladas y retroalimentaci\u00f3n inmediata.",
            BLUE)
        log.append("AZUL: tutor IA asistente agregado")
        break

# ================================================================
# 4. MARCAR ROJO: filas vacias en Tabla 21 y Tabla 22
# ================================================================
# Tabla 21 = doc index 29 (Pruebas Integracion Frontend)
t21 = doc.tables[29]
empty_int = ["AuthFlow.test.tsx", "QuizPage.test.tsx", "QuizPlay.test.tsx",
             "FeedbackPanel.test.tsx", "AssistantChat.test.tsx", "Leaderboard.test.tsx"]
for i, row in enumerate(t21.rows):
    if i == 0: continue
    for e in empty_int:
        if e in row.cells[0].text:
            paint_row(t21, i, RED)
            log.append(f"ROJO Tabla 21: {e}")
            break

# Tabla 22 = doc index 30 (Pruebas Funcionales Frontend)
t22 = doc.tables[30]
for i, row in enumerate(t22.rows):
    if i == 0: continue
    if "session-expiry" in row.cells[0].text:
        paint_row(t22, i, RED)
        log.append("ROJO Tabla 22: session-expiry.spec.ts")
        break

# ================================================================
# 5. AGREGAR RESULTADOS DE LA AUDITORIA
# ================================================================
audit_lines = [
    ("\nResultados de la Auditor\u00eda", False),
    ("", False),
    ("FUNCIONALIDADES VERIFICADAS (BACKEND + FRONTEND):", True),
    ("", False),
    ("1. Autenticaci\u00f3n completa: registro, login (email+JWT), OAuth Google/Microsoft, refresh, logout, forgot/reset password", False),
    ("2. Roles y autorizaci\u00f3n: RolesGuard con STUDENT/TEACHER/ADMIN, decorador @Roles", False),
    ("3. CRUD evaluaciones: crear/editar/eliminar/publicar quizzes con preguntas de 3 tipos", False),
    ("4. Generaci\u00f3n IA de preguntas: QwenProvider + MistralProvider con fallback autom\u00e1tico", False),
    ("5. Resoluci\u00f3n de evaluaciones: StartSession/SubmitAnswer/EndSession con timer y progreso", False),
    ("6. Calificaci\u00f3n autom\u00e1tica: FastAPI + 3 modelos (ExactMatch/Similarity/NLP) + Strategy pattern", False),
    ("7. Retroalimentaci\u00f3n IA: FeedbackPanel + WeaknessChart + AssistantChat + POST /ai/feedback/*", False),
    ("8. Leaderboard global: GET /score/leaderboard con cach\u00e9 Redis TTL 60s", False),
    ("9. Historial puntajes: GET /score/history (propio) + GET /score/history/:userId (docente)", False),
    ("10. Administraci\u00f3n usuarios: CRUD /users + UserAdminPage (activar/desactivar/asignar roles)", False),
    ("11. Perfil p\u00fablico: GET /profile/[username] + ProfilePage con estad\u00edsticas", False),
    ("12. Edici\u00f3n de perfil: PATCH /profile/me + EditProfilePage (nombre, secci\u00f3n)", False),
    ("13. Vista detallada estudiante (docente): /teacher/students/[id]/history con an\u00e1lisis por categor\u00eda", False),
    ("14. Reportes administrador: /admin/reports con estad\u00edsticas globales y gr\u00e1ficos", False),
    ("15. Rate limiting: ThrottlerModule en gateway (10 req/s, 5 req/min auth)", False),
    ("16. Pruebas unitarias: frontend 28 suites, backend 31 suites, IA ~120 pruebas", False),
    ("", False),
    ("FUNCIONALIDADES PARCIALES O NO IMPLEMENTADAS:", True),
    ("", False),
    ("- Logros desbloqueables (HU13): backend parcial (entidad + 4/7 reglas), sin API REST ni frontend", False),
    ("- Notificaciones en tiempo real: entidad existe, implementaci\u00f3n vac\u00eda (SSE controller 0 l\u00edneas)", False),
    ("- Eliminaci\u00f3n de cuenta: DELETE /profile/me existe en backend, sin bot\u00f3n ni flujo frontend", False),
    ("- Store API files (RTK Query): 5 archivos vac\u00edos (baseApi, authApi, quizApi, scoreApi, aiApi)", False),
    ("", False),
    ("HALLAZGOS PRINCIPALES:", True),
    ("", False),
    ("1. Tabla 21 (Pruebas Integraci\u00f3n Frontend): 6 de 7 archivos vac\u00edos. Solo UserAdminSection.test.tsx contiene pruebas.", False),
    ("2. Tabla 22 (Pruebas Funcionales Frontend): session-expiry.spec.ts est\u00e1 vac\u00edo.", False),
    ("3. La afirmaci\u00f3n 'Edici\u00f3n de perfil: p\u00e1gina en construcci\u00f3n' estaba desactualizada. La funcionalidad ya est\u00e1 completa.", False),
    ("4. Perfil p\u00fablico, vista detallada por estudiante, reportes admin, cach\u00e9 leaderboard y rate limiting no est\u00e1n documentados en el cuerpo del documento.", False),
    ("5. Logros y notificaciones requieren desarrollo adicional para alcanzar estado funcional completo.", False),
]

last = doc.paragraphs[-1]
for text, is_bold in audit_lines:
    if text:
        p = add_after(last, text, BLACK, bold=is_bold)
        last = p
    else:
        p = add_after(last, '')
        last = p

log.append("SECCION: Resultados de la Auditor\u00eda agregada al final")

# SAVE
doc.save(SRC)
print(f"Respaldo: {BACKUP}")
print(f"Guardado: {SRC}")
print(f"\nCambios ({len(log)}):")
for l in log:
    print(f"  {l}")

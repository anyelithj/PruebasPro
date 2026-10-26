# -*- coding: utf-8 -*-
"""Fix PURPLE corrections: remove 'HU12 no existe', use only correct info. Add ORANGE observations."""
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"

PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)

doc = Document(SRC)
log = []

def add_para_after(para, text, color):
    el = OxmlElement('w:p')
    para._element.addnext(el)
    np = type(para)(el, doc)
    r = np.add_run(text)
    r.font.color.rgb = color
    return np

# ================================================================
# FIX P552: PURPLE "Gestion de Usuarios" - remove HU12 mention
# ================================================================
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('[CORRECCI') and 'HU12 no existe' in txt and 'gesti' in txt.lower() and 'usuarios' in txt.lower():
        for run in p.runs:
            run.text = ("[CORRECCI\u00d3N] La gesti\u00f3n de usuarios implementa la HU14 (administraci\u00f3n de usuarios: "
                       "listado paginado, b\u00fasqueda por nombre/correo, activar/desactivar cuentas y modificar roles, "
                       "CRUD completo via /admin/users).")
        log.append(f"P{i}: PURPLE corregido - solo HU14")
        add_para_after(p,
            "[OBSERVACI\u00d3N] El texto original asocia err\u00f3neamente esta funcionalidad con HU12. "
            "La retroalimentaci\u00f3n personalizada IA es HU08, ya registrada en el backlog. "
            "La administraci\u00f3n de usuarios corresponde a HU14.",
            ORANGE)
        log.append(f"P{i}: OBSERVACION NARANJA añadida")
        break

# ================================================================
# FIX P556: PURPLE "Gestion Global Evaluaciones" - remove HU12 mention  
# ================================================================
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('[CORRECCI') and 'HU12 no existe' in txt and 'evaluaciones' in txt.lower() and 'global' in txt.lower():
        for run in p.runs:
            run.text = ("[CORRECCI\u00d3N] La gesti\u00f3n global de evaluaciones permite revisar y moderar "
                       "de manera centralizada el cat\u00e1logo de evaluaciones de la plataforma.")
        log.append(f"P{i}: PURPLE corregido - solo moderar evaluaciones")
        add_para_after(p,
            "[OBSERVACI\u00d3N] El texto original asocia err\u00f3neamente esta funcionalidad con HU12. "
            "Esta funci\u00f3n es de moderaci\u00f3n y revisi\u00f3n del cat\u00e1logo, no implementa retroalimentaci\u00f3n IA (HU08).",
            ORANGE)
        log.append(f"P{i}: OBSERVACION NARANJA añadida")
        break

doc.save(SRC)
print(f"Guardado: {SRC}")
print(f"\nCambios ({len(log)}):")
for l in log:
    print(f"  {l}")

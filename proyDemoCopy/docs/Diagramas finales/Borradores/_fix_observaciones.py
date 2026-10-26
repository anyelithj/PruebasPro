# -*- coding: utf-8 -*-
"""Corrige observaciones especificas del usuario sobre prueba.docx."""
import shutil
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"
BACKUP = f"{BASE}/prueba_backup_fixes.docx"
shutil.copy(SRC, BACKUP)

RED    = RGBColor(0xFF, 0x00, 0x00)
BLUE   = RGBColor(0x00, 0x00, 0xFF)
PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
GREEN  = RGBColor(0x00, 0x80, 0x00)
BLACK  = RGBColor(0x00, 0x00, 0x00)

doc = Document(SRC)
log = []

def paint_para(para, color):
    for run in para.runs:
        run.font.color.rgb = color

def paint_run(run, color):
    run.font.color.rgb = color

def replace_run_text(para, old_text, new_text):
    """Replace text in a paragraph's runs."""
    full = para.text
    if old_text not in full:
        return False
    # Find the run containing the old_text and modify it
    for run in para.runs:
        if old_text in run.text:
            run.text = run.text.replace(old_text, new_text)
            return True
    return False

# ================================================================
# FIX 1: CORREGIR TEXTO MORADO P210 - solo entidades completas
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith('[CORRECCI\u00d3N]') and 'ocho entidades' in txt.lower():
        # Replace the run text
        for run in p.runs:
            old = run.text
            if 'ocho entidades' in old or 'Adicionalmente' in old:
                run.text = "Se han determinado las siguientes entidades completamente implementadas en el modelo: Users, AuthSessions, EmailVerificationTokens, PasswordResetTokens, Quizzes, Questions, QuizSessions, Scores y FeedbackIA (persistencia de retroalimentaci\u00f3n)."
                log.append("FIX P210: PURPLE corregido - solo entidades completas")
                break
        break

# ================================================================
# FIX 2: ACTUALIZAR TEXTO NARANJA P211 - incluir ambas entidades
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith('[EXPLICACI\u00d3N T\u00c9CNICA]') and 'Achievement' in txt:
        for run in p.runs:
            if 'Achievement' in run.text:
                run.text = ("Las entidades Achievement y Notification existen como tablas TypeORM pero no est\u00e1n completamente implementadas: "
                           "Achievements carece de API REST y frontend (0 referencias en Frontend/src/); "
                           "Notifications tiene el controlador SSE vac\u00edo (0 l\u00edneas) y casos de uso sin implementar. "
                           "Los diagramas ER y Relacional las excluyen como 'incompletas'.")
                log.append("FIX P211: ORANGE actualizado con ambas entidades")
                break
        break

# ================================================================
# FIX 3: CAMBIAR [NOTA] P212 de VERDE a NARANJA
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith('[NOTA]') and 'Achievement' in txt:
        paint_para(p, ORANGE)
        log.append("FIX P212: [NOTA] Achievement cambiado de VERDE a NARANJA")
        break

# ================================================================
# FIX 4: CORREGIR TEXTO VERDE P550 - HU12 incorrecto
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if 'Gestion de Usuarios' in txt and 'HU12' in txt:
        # Mark the incorrect part red
        for run in p.runs:
            if 'HU12' in run.text or 'retroalimentacion personalizada IA' in run.text:
                paint_run(run, RED)
        log.append("FIX P550: 'HU12' marcado ROJO - referencia incorrecta")
        # Add correction after this paragraph
        el = OxmlElement('w:p')
        p._element.addnext(el)
        np = type(p)(el, doc)
        r = np.add_run("[CORRECCI\u00d3N] La gesti\u00f3n de usuarios implementa la HU14 (administraci\u00f3n de usuarios con CRUD completo). La retroalimentaci\u00f3n personalizada IA corresponde a la HU08, no a HU12.")
        r.font.color.rgb = PURPLE
        log.append("FIX P550: correccion MORADO agregada")
        break

# ================================================================
# FIX 5: CORREGIR TEXTO VERDE P553 - HU12 incorrecto  
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if 'Gestion Global de Evaluaciones' in txt and 'HU12' in txt:
        for run in p.runs:
            if 'HU12' in run.text or 'retroalimentacion personalizada IA' in run.text:
                paint_run(run, RED)
        log.append("FIX P553: 'HU12' marcado ROJO - referencia incorrecta")
        el = OxmlElement('w:p')
        p._element.addnext(el)
        np = type(p)(el, doc)
        r = np.add_run("[CORRECCI\u00d3N] La gesti\u00f3n global de evaluaciones permite revisar y moderar el cat\u00e1logo de evaluaciones. No implementa retroalimentaci\u00f3n personalizada IA (HU08); esa funcionalidad corresponde al m\u00f3dulo de feedback IA (POST /ai/feedback/generate + FeedbackPanel).")
        r.font.color.rgb = PURPLE
        log.append("FIX P553: correccion MORADO agregada")
        break

# ================================================================
# FIX 6: ACTUALIZAR TEXTO AUDITORIA - corregir hallazgo #1
# ================================================================
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith('1) P\u00e1rrafo') and 'ocho entidades' in txt:
        for run in p.runs:
            if 'Correcci' in run.text or 'ocho entidades' in run.text:
                run.text = ("1) P\u00e1rrafo 'ocho entidades... incluye Achievements' marcado ROJO. "
                           "La correcci\u00f3n MORADO lista solo las entidades completamente implementadas "
                           "(Users, AuthSessions, EmailVerificationTokens, PasswordResetTokens, Quizzes, "
                           "Questions, QuizSessions, Scores, FeedbackIA). Achievements y Notification "
                           "se explican en el texto NARANJA como entidades existentes pero incompletas.")
                log.append("FIX: hallazgo #1 en auditoria actualizado")
                break
        break

# ================================================================
# SAVE
# ================================================================
doc.save(SRC)
print(f"Respaldo: {BACKUP}")
print(f"Guardado: {SRC}")
print(f"\nCambios ({len(log)}):")
for l in log:
    print(f"  {l}")

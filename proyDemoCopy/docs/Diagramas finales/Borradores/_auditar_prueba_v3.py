# -*- coding: utf-8 -*-
"""Audita Prueba.docx - enfoque conservador y preciso.
- Solo marca ROJO afirmaciones explicitas INCORRECTAS (funcionalidad que se dice completa pero no lo est\u00e1)
- No marca textos que describan correctamente limitaciones, entidades existentes, o usen tiempo futuro
- Agrega [EXPLICACI\u00d3N T\u00c9CNICA] y [CORRECCI\u00d3N] solo en p\u00e1rrafos con afirmaciones falsas
- Marca filas de tablas de pruebas vac\u00edas como ROJO
- Agrega secci\u00f3n Resultados de la Auditor\u00eda al final
"""
import shutil
from docx import Document
from docx.shared import RGBColor, Pt
from docx.oxml import OxmlElement

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"
BACKUP = f"{BASE}/prueba_backup_auditoria_3.docx"
shutil.copy(SRC, BACKUP)

GREEN  = RGBColor(0x00, 0x80, 0x00)
RED    = RGBColor(0xFF, 0x00, 0x00)
PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
BLUE   = RGBColor(0x00, 0x00, 0xFF)
BLACK  = RGBColor(0x00, 0x00, 0x00)

doc = Document(SRC)
log = []

def add_paragraph_after(para, text, color=None, bold=False, size=None):
    new_p_elem = OxmlElement('w:p')
    para._element.addnext(new_p_elem)
    new_para = type(para)(new_p_elem, doc)
    run = new_para.add_run(text)
    if color:
        run.font.color.rgb = color
    if bold:
        run.font.bold = True
    if size:
        run.font.size = Pt(size)
    return new_para

def paint_para(para, color):
    for run in para.runs:
        run.font.color.rgb = color

def paint_row_red(table, row_idx):
    for cell in table.rows[row_idx].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = RED

def para_txt(para):
    return para.text.strip()

# ============================================================
# 1. PARRAFOS ESPECIFICOS CON AFIRMACIONES INCORRECTAS
# ============================================================
# Buscamos parrafos por texto exacto o patron unico

for i, para in enumerate(doc.paragraphs):
    txt = para_txt(para)
    
    # Skip paragraphs that are already annotations
    if txt.startswith(('[OBSERVACION]', '[NOTA]')):
        continue
        
    # --- CASO 1: Seccion "Notificaciones en tiempo real y centro de notificaciones"
    # Afirma que las notificaciones en tiempo real existen como funcionalidad
    if txt == "Notificaciones en tiempo real y centro de notificaciones":
        paint_para(para, RED)
        log.append(f"ROJO P{i}: '{txt[:60]}...'")
        add_paragraph_after(para, 
            "[EXPLICACI\u00d3N T\u00c9CNICA] El m\u00f3dulo de notificaciones tiene la entidad Notification definida, pero los casos de uso, controlador SSE y gateway WebSocket son archivos vac\u00edos sin implementaci\u00f3n funcional.",
            ORANGE)
        add_paragraph_after(para,
            "[CORRECCI\u00d3N] Notificaciones en tiempo real (propuesta): La entidad Notification existe en el dominio, pero los casos de uso (CreateNotification, GetNotifications, MarkAsRead, DeleteNotification), el controlador SSE y el gateway WebSocket carecen de implementaci\u00f3n. No existe un centro de notificaciones funcional en el frontend.",
            PURPLE)
        continue
        
    # --- CASO 2: Parrafo con afirmacion sobre logros desbloqueables como completos
    # Buscamos texto que diga "logros desbloqueables" como caracteristica implementada
    if ('logros desbloqueables' in txt.lower() or 'logros' in txt.lower()) and \
       not any(w in txt.lower() for w in ['parcial', 'pendiente', 'no tiene', 'no hay', 'falta',
                                           'sin api', 'sin frontend', 'sin integraci', 'CERO',
                                           'no cumpl', 'faltante', 'no existe', 'backend parcial',
                                           'HU13']) and \
       any(w in txt.lower() for w in ['completamente', 'implementad', 'funcionalidad completa',
                                       'totalmente', 'integrad', 'disponible', 'funcional', 'permite']):
        # Esta verificando si el parrafo AFIRMA que los logros estan implementados
        paint_para(para, RED)
        log.append(f"ROJO P{i}: '{txt[:80]}...'")
        add_paragraph_after(para,
            "[EXPLICACI\u00d3N T\u00c9CNICA] Los logros desbloqueables existen solo como backend parcial: entidad Achievement, repositorio y 4/7 reglas (FIRST_QUIZ, PERFECT_SCORE, POINTS_1000, POINTS_5000). No hay endpoint REST ni componente frontend.",
            ORANGE)
        add_paragraph_after(para,
            "[CORRECCI\u00d3N] Los logros desbloqueables (HU13) se encuentran parcialmente implementados en el backend (entidad Achievement, repositorio, 4 reglas de desbloqueo v\u00eda CheckAchievementsUseCase). Est\u00e1n pendientes: endpoint REST GET /achievements, frontend (componentes, p\u00e1ginas, rutas API, store) e integraci\u00f3n backend-frontend.",
            PURPLE)
        continue
        
    # --- CASO 3: "7 suites de integración" como frase que afirma completitud
    if '7 suite' in txt.lower() and 'integraci' in txt.lower() and \
       not any(w in txt.lower() for w in ['vacio', 'vací', 'pendiente', 'no ejecutad', '6 de']):
        paint_para(para, RED)
        log.append(f"ROJO P{i}: '{txt[:80]}...'")
        add_paragraph_after(para,
            "[EXPLICACI\u00d3N T\u00c9CNICA] De las 7 suites de pruebas de integraci\u00f3n del frontend, 6 archivos est\u00e1n vac\u00edos (0 bytes): AuthFlow.test.tsx, QuizPage.test.tsx, QuizPlay.test.tsx, FeedbackPanel.test.tsx, AssistantChat.test.tsx, Leaderboard.test.tsx.",
            ORANGE)
        add_paragraph_after(para,
            "[CORRECCI\u00d3N] El frontend cuenta con 7 suites de integraci\u00f3n planificadas, de las cuales solo 1 (api.service.spec.ts) contiene pruebas implementadas. Las 6 restantes son archivos vac\u00edos sin casos de prueba.",
            PURPLE)

    # --- CASO 4: "eliminación de cuenta" como funcionalidad completa
    if any(w in txt.lower() for w in ['eliminar cuenta', 'eliminar mi cuenta', 'deleteAccount', 'eliminaci\u00f3n de cuenta']) and \
       any(w in txt.lower() for w in ['disponible', 'permite', 'implement', 'complet', 'funcionalidad', 'puede']) and \
       not any(w in txt.lower() for w in ['solo backend', 'pendiente', 'sin frontend', 'backend parcial', 'falta', 'no tiene']):
        paint_para(para, RED)
        log.append(f"ROJO P{i}: '{txt[:80]}...'")
        add_paragraph_after(para,
            "[EXPLICACI\u00d3N T\u00c9CNICA] El endpoint DELETE /profile/me existe en el backend, pero no hay ning\u00fan componente frontend (bot\u00f3n, formulario, flujo de confirmaci\u00f3n) que implemente la eliminaci\u00f3n de cuenta.",
            ORANGE)
        add_paragraph_after(para,
            "[CORRECCI\u00d3N] La eliminaci\u00f3n de cuenta est\u00e1 implementada \u00fanicamente en el backend a trav\u00e9s del endpoint DELETE /profile/me. El frontend carece de cualquier interfaz de usuario para esta funcionalidad, por lo que no est\u00e1 disponible para los usuarios finales.",
            PURPLE)

# ============================================================
# 2. TABLAS DE PRUEBAS VACIAS -> ROJO
# ============================================================
t29 = doc.tables[29]
empty_fe_int = ["AuthFlow.test.tsx", "QuizPage.test.tsx", "QuizPlay.test.tsx",
                "FeedbackPanel.test.tsx", "AssistantChat.test.tsx", "Leaderboard.test.tsx"]
for i, row in enumerate(t29.rows):
    if i == 0:
        continue
    for e in empty_fe_int:
        if e in row.cells[0].text:
            paint_row_red(t29, i)
            log.append(f"ROJO T29 fila {i}: {e}")
            break

t30 = doc.tables[30]
for i, row in enumerate(t30.rows):
    if i == 0:
        continue
    if "session-expiry" in row.cells[0].text:
        paint_row_red(t30, i)
        log.append(f"ROJO T30 fila {i}: session-expiry.spec.ts")
        break

# ============================================================
# 3. SECCION "Resultados de la Auditoría"
# ============================================================
audit_lines = [
    ("Resultados de la Auditor\u00eda", True),
    ("", False),
    ("1. Funcionalidades completamente verificadas (backend + frontend + integraci\u00f3n):", True),
    ("", False),
    ("- Registro de usuarios: RegisterUseCase + /auth/register + RegisterForm + verificaci\u00f3n email", False),
    ("- Inicio de sesi\u00f3n: LoginUseCase + /auth/login + LoginForm + OAuth Google/Microsoft", False),
    ("- Recuperaci\u00f3n de contrase\u00f1a: RequestPasswordResetUseCase + ResetPasswordUseCase + reenv\u00edo", False),
    ("- Creaci\u00f3n de evaluaciones: CRUD quizzes (crear/editar/eliminar) con roles TEACHER+", False),
    ("- Generaci\u00f3n de preguntas con IA: QwenProvider + MistralProvider con fallback autom\u00e1tico", False),
    ("- Resoluci\u00f3n de evaluaciones: StartSession / SubmitAnswer / EndSession + calificaci\u00f3n autom\u00e1tica", False),
    ("- Leaderboard global: GET /score/leaderboard con cach\u00e9 Redis TTL 60s", False),
    ("- Historial de puntajes: GET /score/history (propio) y GET /score/history/:userId (p\u00fablico)", False),
    ("- Administraci\u00f3n de usuarios: CRUD completo via /users para rol ADMIN", False),
    ("- Perfil p\u00fablico: GET /profile/[username] con estad\u00edsticas", False),
    ("- Edici\u00f3n de perfil: PATCH /profile/me con avatar, biography, social links", False),
    ("- Retroalimentaci\u00f3n IA: FeedbackPanel + WeaknessChart + AssistantChat (frontend)", False),
    ("- Persistencia de feedback_ia: SaveFeedbackUseCase + SubmitAnswerFlow", False),
    ("- Rate limiting global: ThrottlerModule configurado en gateway API", False),
    ("- Pruebas unitarias frontend: 28 suites, ~170 pruebas", False),
    ("- Pruebas unitarias backend: 31+ suites, ~152 pruebas", False),
    ("- Pruebas unitarias microservicio IA: ~120 pruebas", False),
    ("", False),
    ("2. Funcionalidades documentadas pero NO implementadas (o solo parcialmente):", True),
    ("", False),
    ("- Logros desbloqueables (HU13): backend parcial (4/7 reglas, sin endpoint REST, sin frontend, sin integraci\u00f3n)", False),
    ("- Pruebas de integraci\u00f3n frontend: 6 de 7 archivos vac\u00edos (AuthFlow, QuizPage, QuizPlay, FeedbackPanel, AssistantChat, Leaderboard)", False),
    ("- Pruebas funcionales frontend: 1 archivo vac\u00edo (session-expiry)", False),
    ("- Eliminaci\u00f3n de cuenta: DELETE /profile/me en backend, sin interfaz de usuario frontend", False),
    ("- Notificaciones en tiempo real: m\u00f3dulo con entidad pero sin implementaci\u00f3n funcional en backend ni frontend", False),
    ("- Pruebas de rendimiento/carga: planificadas en el plan de pruebas pero no ejecutadas", False),
    ("", False),
    ("3. Funcionalidades implementadas pero NO documentadas:", True),
    ("", False),
    ("- Rate limiting global: ThrottlerModule en gateway (no mencionado en el documento)", False),
    ("- Cach\u00e9 de leaderboard: CacheModule TTL 60s con Redis (no mencionado)", False),
    ("- Feedback_ia en historial: SaveFeedbackUseCase integrado en SubmitAnswerFlow (no documentado)", False),
    ("", False),
    ("4. Hallazgos principales:", True),
    ("", False),
    ("1) 7 suites de integraci\u00f3n en frontend: 6 archivos est\u00e1n vac\u00edos. Solo api.service.spec.ts contiene pruebas.", False),
    ("2) Logros desbloqueables: backend parcial sin API REST ni frontend. No debieran listarse como completos.", False),
    ("3) Eliminaci\u00f3n de cuenta: endpoint backend existe pero no hay interfaz de usuario.", False),
    ("4) Notificaciones en tiempo real: referenciadas como funcionalidad existente pero el m\u00f3dulo est\u00e1 vac\u00edo.", False),
    ("5) El diccionario de datos incluye tabla achievements correctamente (la entidad existe), pero el documento la presenta como funcionalidad completa.", False),
    ("6) No se documentan: rate limiting, cach\u00e9 Redis, ni persistencia de feedback_ia en historial.", False),
]

last_para = doc.paragraphs[-1]
for text, is_bold in audit_lines:
    if text:
        p = add_paragraph_after(last_para, text, BLACK, bold=is_bold)
        last_para = p
    else:
        p = add_paragraph_after(last_para, '')
        last_para = p

log.append("SECCI\u00d3N: Resultados de la Auditor\u00eda agregada al final")

# ============================================================
# 4. SAVE
# ============================================================
doc.save(SRC)
print(f"Respaldo: {BACKUP}")
print(f"Guardado: {SRC}")
print(f"\nCambios realizados ({len(log)}):")
for l in log:
    print(f"  {l}")

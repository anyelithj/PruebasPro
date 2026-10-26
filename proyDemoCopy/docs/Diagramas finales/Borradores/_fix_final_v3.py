# -*- coding: utf-8 -*-
"""Fixes: P304 (quitar referencias incompletas), P550/P554 (HU12 incorrecto → ROJO + MORADO)."""
from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement

BASE = "C:/Users/angui/Documents/proyectoBorrador/proyDemo/Diagramas finales"
SRC = f"{BASE}/prueba.docx"

RED    = RGBColor(0xFF, 0x00, 0x00)
PURPLE = RGBColor(0x80, 0x00, 0x80)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)

doc = Document(SRC)
log = []

def paint_para(para, color):
    for run in para.runs:
        run.font.color.rgb = color

def add_para_after(para, text, color):
    el = OxmlElement('w:p')
    para._element.addnext(el)
    np = type(para)(el, doc)
    r = np.add_run(text)
    r.font.color.rgb = color
    return np

# ================================================================
# FIX 1: P304 - quitar "idiomas", "tipos de notificacion", "codigos de logro"
# ================================================================
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('En la capa de dominio') and p.runs and str(p.runs[0].font.color.rgb) == '7030A0':
        old = p.text
        # Remove the problematic references from the text
        new_text = old.replace(', idiomas', '')
        new_text = new_text.replace('idiomas, ', '')
        new_text = new_text.replace(', tipos de notificacion', '')
        new_text = new_text.replace('tipos de notificacion, ', '')
        new_text = new_text.replace(', codigos de logro', '')
        new_text = new_text.replace('codigos de logro y ', '')  # Edge case
        if 'codigos de logro' in new_text:
            new_text = new_text.replace('codigos de logro', '')
        # Clean up any double commas or spacing
        new_text = new_text.replace(', ,', ',')
        new_text = new_text.replace('  ', ' ')
        new_text = new_text.strip()
        for run in p.runs:
            run.text = new_text
        log.append(f"P304: quitados idiomas, tipos de notificacion, codigos de logro")
        
        # Add ORANGE explanation after
        add_para_after(p,
            "[EXPLICACI\u00d3N T\u00c9CNICA] Los enumerados 'idiomas', 'tipos de notificaci\u00f3n' y 'c\u00f3digos de logro' existen como definitiones en el c\u00f3digo fuente "
            "(Language enum en quiz/domain, NotificationType enum en notifications/domain, AchievementCode enum en score/domain) pero las funcionalidades asociadas "
            "no est\u00e1n completas end-to-end: idiomas solo tiene implementado ESPA\u00d1OL (sin selector de idioma ni i18n en frontend), "
            "NotificationType tiene el m\u00f3dulo de notificaciones con controlador SSE vac\u00edo (0 l\u00edneas), "
            "AchievementCode tiene backend completo con 4 reglas de desbloqueo pero 0 referencias en Frontend/src/.",
            ORANGE)
        log.append("P304: añadida explicacion NARANJA sobre items quitados")
        break

# ================================================================
# FIX 2: P550 (Gestion de Usuarios) - marcar ROJO, convertir P551 en MORADO
# ================================================================
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Gestion de Usuarios') and 'HU12' in txt:
        # Mark as RED (contains incorrect info)
        paint_para(p, RED)
        log.append(f"P550: marcado ROJO (contiene referencia incorrecta a HU12)")
        
        # Find the next paragraph (P551 orange observation) and replace its text with PURPLE correction
        # Look at the next paragraph
        next_p = doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else None
        if next_p and '[OBSERVACI' in next_p.text:
            paint_para(next_p, PURPLE)
            for run in next_p.runs:
                run.text = ("[CORRECCI\u00d3N] La gesti\u00f3n de usuarios permite ver listado paginado, "
                           "buscar por nombre/correo, activar/desactivar cuentas y modificar roles. "
                           "HU12 no existe en el backlog del producto (solo HU01-HU11 y HU14). "
                           "Esta funcionalidad implementa la HU14 (administraci\u00f3n de usuarios con CRUD completo via /admin/users). "
                           "La retroalimentaci\u00f3n personalizada IA corresponde a HU08, no a HU12.")
            log.append(f"P551: convertido de OBSERVACION a CORRECCION MORADO")
        break

# ================================================================
# FIX 3: P554 (Gestion Global de Evaluaciones) - marcar ROJO, convertir P555 en MORADO
# ================================================================
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith('Gestion Global de Evaluaciones') and 'HU12' in txt:
        paint_para(p, RED)
        log.append(f"P554: marcado ROJO (contiene referencia incorrecta a HU12)")
        
        next_p = doc.paragraphs[i + 1] if i + 1 < len(doc.paragraphs) else None
        if next_p and '[OBSERVACI' in next_p.text:
            paint_para(next_p, PURPLE)
            for run in next_p.runs:
                run.text = ("[CORRECCI\u00d3N] La gesti\u00f3n global de evaluaciones permite revisar y moderar "
                           "el cat\u00e1logo de evaluaciones de la plataforma. "
                           "HU12 no existe en el backlog del producto. "
                           "La retroalimentaci\u00f3n personalizada IA corresponde a HU08 (POST /ai/feedback/generate + FeedbackPanel), "
                           "no a esta funcionalidad.")
            log.append(f"P555: convertido de OBSERVACION a CORRECCION MORADO")
        break

# ================================================================
# SAVE
# ================================================================
doc.save(SRC)
print(f"Guardado: {SRC}")
print(f"\nCambios ({len(log)}):")
for l in log:
    print(f"  {l}")

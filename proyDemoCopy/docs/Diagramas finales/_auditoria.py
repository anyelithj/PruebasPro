# -*- coding: utf-8 -*-
"""Auditoria de prueba.docx contra el codigo fuente de PROYDEMO.
Aplica leyenda de colores: Verde/Rojo/Morado/Naranja/Azul.
Trabaja sobre prueba.docx (respaldo previo: prueba_backup_preauditoria.docx)."""
from docx import Document
from docx.shared import RGBColor, Pt

SRC = "prueba.docx"

VERDE   = RGBColor(0x00, 0xB0, 0x50)
ROJO    = RGBColor(0xFF, 0x00, 0x00)
MORADO  = RGBColor(0x70, 0x30, 0xA0)
NARANJA = RGBColor(0xED, 0x7D, 0x31)
AZUL    = RGBColor(0x00, 0x70, 0xC0)

doc = Document(SRC)
log = []

# ---------- helpers ----------
def color_cell(cell, rgb):
    for p in cell.paragraphs:
        if not p.runs and p.text:
            # rebuild single run
            r = p.add_run(p.text)
        for run in p.runs:
            run.font.color.rgb = rgb

def add_cell_para(cell, text, rgb, bold=False, size=9):
    p = cell.add_paragraph()
    r = p.add_run(text)
    r.font.color.rgb = rgb
    r.bold = bold
    r.font.size = Pt(size)
    return p

def add_row(table, values, rgb):
    row = table.add_row()
    for i, v in enumerate(values):
        if i < len(row.cells):
            c = row.cells[i]
            c.text = v
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = rgb
    return row

def find_para_index(needle):
    for i, p in enumerate(doc.paragraphs):
        if needle in p.text:
            return i
    return -1

def insert_note_before(idx, text, rgb, bold=False):
    anchor = doc.paragraphs[idx]
    newp = anchor.insert_paragraph_before()
    r = newp.add_run(text)
    r.font.color.rgb = rgb
    r.bold = bold
    return newp

T = doc.tables

# ==========================================================
# 1) TABLAS DE RESULTADOS DE PRUEBAS
# ==========================================================
# T29 Frontend integracion: archivos VACIOS -> ROJO (sin evidencia)
EMPTY_FE_INT = ["AuthFlow.test.tsx","QuizPage.test.tsx","QuizPlay.test.tsx",
                "FeedbackPanel.test.tsx","AssistantChat.test.tsx","Leaderboard.test.tsx"]
t = T[29]
for row in t.rows[1:]:
    txt = row.cells[0].text.strip()
    if any(n in txt for n in EMPTY_FE_INT):
        for c in row.cells: color_cell(c, ROJO)
        log.append(f"ROJO T29: {txt}")
# nota naranja en celda del primer vacio
add_cell_para(t.rows[2].cells[1], "Auditoria: archivo de prueba vacio (0 bytes); sin evidencia de ejecucion.", NARANJA)

# T30 Frontend funcional: solo auth.spec.ts vacio -> ROJO
t = T[30]
for row in t.rows[1:]:
    txt = row.cells[0].text.strip()
    if txt.endswith("functional/auth.spec.ts"):
        for c in row.cells: color_cell(c, ROJO)
        add_cell_para(row.cells[1], "Auditoria: archivo vacio (0 bytes); prueba no implementada.", NARANJA)
        log.append(f"ROJO T30: {txt}")

# T31 Backend unitario: falta app.controller.spec.ts (real) -> AZUL
t = T[31]
existing = {r.cells[0].text.strip() for r in t.rows}
if not any("app.controller.spec.ts" in e for e in existing):
    add_row(t, ["src/app.controller.spec.ts", "Controlador principal de la aplicacion (1 caso, real)"], AZUL)
    log.append("AZUL T31 (+fila): src/app.controller.spec.ts")

# T32 Backend integracion: falta notifications/sse (real, 6 casos) -> AZUL
t = T[32]
existing = {r.cells[0].text.strip() for r in t.rows}
if not any("sse.integration" in e for e in existing):
    add_row(t, ["tests/integration/notifications/sse.integration.spec.ts", "Canal SSE de notificaciones (6 casos, real)"], AZUL)
    log.append("AZUL T32 (+fila): notifications/sse.integration.spec.ts")

# ==========================================================
# 2) REQUERIMIENTOS FUNCIONALES (T1, col 2 = REQUERIMIENTO)
# ==========================================================
t = T[1]
REQ_COL = 2
# R1..R13 (filas 1..13) verificados implementados -> VERDE
for ri in range(1, 14):
    color_cell(t.rows[ri].cells[REQ_COL], VERDE)
# Fila 14 (#14) "global y por categoria": parte "por categoria" NO implementada
r14 = t.rows[14].cells[REQ_COL]
color_cell(r14, ROJO)
add_cell_para(r14, "Correccion: El producto cuenta con una clasificacion global de mejores puntajes, "
                   "filtrable por periodo y por seccion, actualizada con cache de 60 segundos. "
                   "La clasificacion por categoria tematica no esta implementada.", MORADO)
add_cell_para(r14, "Explicacion: GetLeaderboardUseCase y score.repository.getLeaderboard aceptan los "
                   "parametros limit, months (periodo) y section; no existe filtro por categoria "
                   "tematica del quiz. Ademas este requerimiento duplica al #13.", NARANJA)
log.append("VERDE T1 R1-13; ROJO+MORADO+NARANJA T1 R14 (por categoria)")
# AZUL: requerimientos implementados ausentes en el documento
add_row(t, ["15", "Administracion",
            "El sistema debe permitir al administrador gestionar los usuarios (listar, crear, editar, "
            "activar/desactivar y asignar roles) para controlar el acceso a la plataforma. "
            "[Implementado: users.controller.ts]"], AZUL)
add_row(t, ["16", "Gamificacion",
            "El sistema desbloquea y persiste logros del estudiante al completar una evaluacion "
            "(primer quiz, puntaje perfecto, umbrales de puntos). NOTA: implementado solo en backend "
            "(persistencia); sin API/notificacion/UI. Ver Observaciones de Auditoria."], AZUL)
log.append("AZUL T1: +Administracion (#15), +Gamificacion/Logros (#16)")

# ==========================================================
# 3) REQUERIMIENTOS NO FUNCIONALES (T2, col 2)
# ==========================================================
t = T[2]
NF_COL = 2
for ri in range(1, 8):
    color_cell(t.rows[ri].cells[NF_COL], VERDE)
# R2 Rendimiento: metas no verificadas (pruebas de rendimiento vacias)
add_cell_para(t.rows[2].cells[NF_COL],
    "Explicacion de auditoria: las metas de rendimiento no fueron verificadas; los archivos de "
    "pruebas de rendimiento (performance/*.spec.ts y scenarios/*.py) estan vacios (0 bytes).", NARANJA)
# R6 Mantenibilidad: cobertura >80% no verificable
add_cell_para(t.rows[6].cells[NF_COL],
    "Explicacion de auditoria: la cobertura >80% no se pudo verificar; existen multiples archivos de "
    "prueba vacios (integracion frontend, funcionales y de usabilidad/rendimiento). La documentacion "
    "OpenAPI/Swagger y los comentarios de patrones si estan presentes.", NARANJA)
# R7 Accesibilidad: parcial
add_cell_para(t.rows[7].cells[NF_COL],
    "Explicacion de auditoria: la accesibilidad se valida parcialmente con jest-axe en atomos y "
    "moleculas; navegacion por teclado, contraste y compatibilidad entre navegadores estan "
    "planificadas (archivos de prueba vacios).", NARANJA)
log.append("VERDE T2 R1-7; NARANJA notas en R2/R6/R7")
# AZUL: NFR implementado ausente
add_row(t, ["8", "Seguridad/Disponibilidad",
            "El backend implementa limitacion de tasa (rate limiting) global mediante ThrottlerModule "
            "y cache-aside en la clasificacion. [Implementado: app.module.ts, get-leaderboard.use-case.ts]"], AZUL)
log.append("AZUL T2: +rate limiting/cache (#8)")

# ==========================================================
# 4) HISTORIAS DE USUARIO (T3, col 1 = HISTORIA)
# ==========================================================
t = T[3]
HU_COL = 1
for ri in range(1, len(t.rows)):
    color_cell(t.rows[ri].cells[HU_COL], VERDE)
# HU11 (fila 11) "diferentes categorias": ranking por categoria no implementado
hu11 = t.rows[11].cells[HU_COL]
color_cell(hu11, ROJO)
add_cell_para(hu11, "Correccion: Como estudiante, quiero ver mi posicion en la clasificacion general "
                    "(filtrable por periodo y seccion) para comparar mi desempeno y mantener mi motivacion.", MORADO)
add_cell_para(hu11, "Explicacion: la clasificacion global esta implementada (leaderboard con cache 60s), "
                    "pero el filtro por categoria tematica no existe en el codigo (solo limit/periodo/seccion).", NARANJA)
# HU10 nota: dashboard docente parcial
add_cell_para(t.rows[10].cells[HU_COL],
    "Explicacion de auditoria: existe la vista de historial por estudiante (teacher/students/[id]/history); "
    "el dashboard docente con distribucion y promedios esta en fase de implementacion (ver doc.).", NARANJA)
log.append("VERDE T3; ROJO+MORADO+NARANJA HU11; NARANJA HU10")

# ==========================================================
# 5) LOGROS: referencia en INTRODUCCION (AZUL) + nota en P377 (NARANJA)
# ==========================================================
idx = find_para_index("Descripci")  # 'Descripcion de la Empresa'
# buscar exactamente el encabezado de empresa
emp = -1
for i,p in enumerate(doc.paragraphs):
    if "Empresa" in p.text and "Descrip" in p.text:
        emp = i; break
if emp != -1:
    insert_note_before(emp,
        "El sistema incorpora ademas un modulo de Logros Desbloqueables (Achievements) que registra "
        "automaticamente hitos del estudiante (primer quiz, puntaje perfecto y umbrales de puntos) al "
        "completar una evaluacion. Nota de auditoria: este modulo esta implementado en la capa de "
        "persistencia del backend; su exposicion via API, notificaciones e interfaz de usuario esta "
        "pendiente (ver seccion Observaciones y Explicaciones de Auditoria).", AZUL)
    log.append("AZUL intro: referencia a modulo de Logros")

# nota naranja tras el paquete Historial/Retroalimentacion (acceso a logros)
pk = -1
for i,p in enumerate(doc.paragraphs):
    if "Este paquete" in p.text and "seguimiento del desempeno" in p.text.replace("ñ","n"):
        pk = i; break
if pk == -1:
    for i,p in enumerate(doc.paragraphs):
        if "acceso al historial de puntajes y logros" in p.text:
            pk = i; break
if pk != -1 and pk+1 < len(doc.paragraphs):
    insert_note_before(pk+1,
        "Nota de auditoria (Naranja): la verificacion y persistencia de logros si esta implementada en el "
        "backend, pero su VISUALIZACION por el usuario (acceso a logros desde la interfaz) NO esta "
        "implementada; no existe endpoint REST ni UI de logros.", NARANJA)
    log.append("NARANJA nota tras paquete Historial (acceso a logros)")

doc.save(SRC)
print("\n".join(log))
print(f"\nTotal cambios: {len(log)}")
